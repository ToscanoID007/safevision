#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Reporte tecnico final (docs/reporte-final.md) con el mismo diseno institucional
que los informes de servicio social: plantilla_itc.docx (encabezado y pie del
ITC), titulos en negrita cursiva, tablas azules, tabla de ilustraciones con
miniaturas y pies 'Fig N.'. Arial 12, interlineado 1,5, texto justificado.

    python3 generar_reporte_tecnico.py [paginas.json]

Convierte el Markdown del reporte: titulos ##/###/####, parrafos, listas,
tablas, notas (>), diagramas mermaid (sustituidos por docs/entrega/img/
diagrama-N.png) y la tabla de figuras de evidencias. Requiere python-docx y Pillow.
"""
import json
import os
import re
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

import generar_informe_equipo as base   # mismo diseno: utilidades y plantilla

AQUI = base.AQUI
RAIZ = base.RAIZ
MD = os.path.join(RAIZ, "docs", "reporte-final.md")
SALIDA = os.path.join(RAIZ, "docs", "entrega", "reporte-final.docx")
IMG = base.IMG

DIAGRAMAS = [
    "Arquitectura en tres capas: la PC y el robot, con el Robot Server como única frontera HTTP ↔ ROS.",
    "Opción A: la profundidad de la Astra Pro como segunda fuente del costmap, sin tocar /scan.",
]
SIMBOLOS = {r"\theta": "θ", r"\omega": "ω"}
ORDEN_FIG = []   # [(ruta, descripcion)]


# ---------------------------------------------------------------- texto en linea
def limpiar(t):
    return re.sub(r"(✅|❌|⚠️|📝)\s*", "", t)


def en_linea(par, t, tam=12, negrita=False, cursiva=False, color=None):
    """**negrita**, *cursiva*, `codigo`, [texto](url), <url>, $math$."""
    t = limpiar(t)
    patron = r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|<https?:[^>]+>|\$[^$]+\$|~~[^~]+~~)"
    for trozo in re.split(patron, t):
        if not trozo:
            continue
        if trozo.startswith("**"):
            en_linea(par, trozo[2:-2], tam, True, cursiva, color)
        elif trozo.startswith("~~"):
            r = par.add_run(trozo[2:-2]); base.fuente(r, tam, negrita, cursiva, color); r.font.strike = True
        elif trozo.startswith("*"):
            en_linea(par, trozo[1:-1], tam, negrita, True, color)
        elif trozo.startswith("`"):
            r = par.add_run(trozo[1:-1]); base.fuente(r, tam - 1.5, negrita, cursiva, RGBColor(0x33, 0x33, 0x33))
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas"); r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        elif trozo.startswith("$"):
            m = trozo[1:-1]
            for a, b in SIMBOLOS.items():
                m = m.replace(a, b)
            base.fuente(par.add_run(m), tam, negrita, True, color)
        elif trozo.startswith("<"):
            base.fuente(par.add_run(trozo[1:-1]), tam, negrita, cursiva, RGBColor(0x11, 0x55, 0xCC))
        elif trozo.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", trozo)
            etiqueta, destino = m.group(1), m.group(2)
            if destino.startswith("http"):
                en_linea(par, etiqueta, tam, negrita, cursiva, RGBColor(0x11, 0x55, 0xCC))
            else:
                en_linea(par, etiqueta, tam, negrita, cursiva, color)
        else:
            base.fuente(par.add_run(trozo), tam, negrita, cursiva, color)


# ---------------------------------------------------------------- bloques
def tabla_md(doc, filas):
    cab, cuerpo = filas[0], filas[1:]
    n = len(cab)
    largos = [4] * n
    for f in filas:
        for i, c in enumerate(f[:n]):
            largos[i] = max(largos[i], min(len(re.sub(r"[*`]", "", c)), 60))
    total = sum(largos)
    ancho_util = 15.5
    anchos = [max(1.6, ancho_util * l / total) for l in largos]
    k = ancho_util / sum(anchos)
    anchos = [a * k for a in anchos]
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cab):
        cel = t.rows[0].cells[i]; cel.text = ""
        p = cel.paragraphs[0]; base.formato(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 2, 2)
        en_linea(p, c, 10, True, True, RGBColor(0xFF, 0xFF, 0xFF))
        base.sombra(cel, base.AZUL_CABECERA)
    for fila in cuerpo:
        celdas = t.add_row().cells
        for i in range(n):
            c = fila[i] if i < len(fila) else ""
            celdas[i].text = ""
            p = celdas[i].paragraphs[0]; base.formato(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 2, 2)
            en_linea(p, c, 10)
            base.sombra(celdas[i], base.AZUL_CLARO)
    t.autofit = False
    for i, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(anchos[i] * 567)))
    for fila in t.rows:
        for i, cel in enumerate(fila.cells):
            cel.width = Cm(anchos[i])
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + lado)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), base.AZUL_OSCURO)
        bordes.append(e)
    t._tbl.tblPr.append(bordes)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figura(doc, ruta, desc, ancho=13.0):
    ORDEN_FIG.append((ruta, desc))
    n = len(ORDEN_FIG)
    w, h = Image.open(ruta).size
    if ancho * h / w > 10.5:
        ancho = 10.5 * w / h
    p = doc.add_paragraph(); base.formato(p, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 8, 2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(ruta, width=Cm(ancho))
    cap = doc.add_paragraph(); base.formato(cap, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 2, 12)
    en_linea(cap, "Fig %d. %s" % (n, desc), 10, cursiva=True)


def nota(doc, texto_nota):
    p = doc.add_paragraph(); base.formato(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.3, 6, 10)
    p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.right_indent = Cm(0.4)
    pPr = p._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); izq = OxmlElement("w:left")
    izq.set(qn("w:val"), "single"); izq.set(qn("w:sz"), "18"); izq.set(qn("w:space"), "8"); izq.set(qn("w:color"), base.AZUL_CABECERA)
    bdr.append(izq); pPr.append(bdr)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EAF1FB")
    pPr.append(shd)
    en_linea(p, texto_nota, 11)


def ruta_imagen(rel):
    r = os.path.join(RAIZ, "docs", rel)
    if r.endswith(".webp"):
        r = os.path.join(IMG, os.path.basename(r)[:-5] + ".png")
    return r


# ---------------------------------------------------------------- conversion
def convertir(doc, md, vinetas, numeradas):
    lineas = md.split("\n")
    i = next(k for k, l in enumerate(lineas) if l.startswith("## 1. Datos generales"))
    titulos, diagrama = [], 0
    primer_h1 = True
    while i < len(lineas):
        l = lineas[i]
        if not l.strip() or re.match(r"^---\s*$", l):
            i += 1; continue
        m = re.match(r"^(#{2,4})\s+(.*)$", l)
        if m:
            nivel = len(m.group(1)) - 1
            t = limpiar(m.group(2)).replace("`", "")
            if nivel == 1:
                if not primer_h1:
                    base.salto(doc)
                primer_h1 = False
                t = t.rstrip(".") + "."
            base.titulo(doc, t, nivel)
            if nivel <= 2:
                titulos.append((nivel, t))
            if t == "Firmas.":
                i += 1
                while i < len(lineas) and not lineas[i].startswith("|"):
                    if lineas[i].strip():
                        p = doc.add_paragraph(); base.formato(p, WD_ALIGN_PARAGRAPH.LEFT, 1.5, 6, 6)
                        en_linea(p, lineas[i].strip())
                    i += 1
                while i < len(lineas) and lineas[i].startswith("|"):
                    i += 1
                firmas(doc)
                continue
            i += 1; continue
        if l.startswith("```"):
            lang = l[3:].strip()
            j = i + 1
            while j < len(lineas) and not lineas[j].startswith("```"):
                j += 1
            if lang == "mermaid":
                diagrama += 1
                figura(doc, os.path.join(IMG, "diagrama-%d.png" % diagrama), DIAGRAMAS[diagrama - 1],
                       9 if diagrama == 1 else 15)
            i = j + 1; continue
        if l.startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].startswith("|"):
                bloque.append(lineas[i]); i += 1
            filas = [[c.strip() for c in b.strip().strip("|").split("|")] for b in bloque
                     if not re.match(r"^\|\s*:?-{2,}", b)]
            if any("![" in b for b in bloque):
                for f in range(len(filas)):
                    for k, c in enumerate(filas[f]):
                        im = re.search(r"!\[[^\]]*\]\(([^)]+)\)", c)
                        if im:
                            pie = (filas[f + 1][k] if f + 1 < len(filas) else "").strip("*")
                            pie = re.sub(r"^Fig\.\s*\d+\.\s*", "", pie)
                            figura(doc, ruta_imagen(im.group(1)), pie)
                continue
            if filas and all(c == "" for c in filas[0]):
                filas = filas[1:]
            tabla_md(doc, filas)
            continue
        if l.startswith(">"):
            bloque = []
            while i < len(lineas) and lineas[i].startswith(">"):
                bloque.append(lineas[i].lstrip(">").strip()); i += 1
            nota(doc, " ".join(bloque))
            continue
        m = re.match(r"^(\s*)(-|\d+\.)\s+(.*)$", l)
        if m:
            nid = numeradas() if m.group(2) != "-" else vinetas()
            while i < len(lineas):
                mm = re.match(r"^(\s*)(-|\d+\.)\s+(.*)$", lineas[i])
                if not mm:
                    break
                t = mm.group(3); i += 1
                while i < len(lineas) and re.match(r"^\s{2,}\S", lineas[i]) and not re.match(r"^\s*(-|\d+\.)\s", lineas[i]):
                    t += " " + lineas[i].strip(); i += 1
                p = doc.add_paragraph()
                base.formato(p, WD_ALIGN_PARAGRAPH.LEFT if "http" in t else WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 0, 4)
                base.numerar(p, nid); en_linea(p, t)
            continue
        texto_par = []
        while i < len(lineas) and lineas[i].strip() and not re.match(r"^(#|```|\||>|\s*(-|\d+\.)\s)", lineas[i]) \
                and not re.match(r"^---\s*$", lineas[i]):
            texto_par.append(lineas[i].strip()); i += 1
        p = doc.add_paragraph(); base.formato(p)
        unido = " ".join(texto_par)
        if re.search(r":\**\s*$", unido):
            p.paragraph_format.keep_with_next = True
        en_linea(p, unido)
    return titulos


def firmas(doc):
    t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    datos = [("Luis Adrian Flores Bueno", "Prestador · No. de control 22460736"),
             ("Andros Jair Toscano Farias", "Prestador · No. de control 22460548"),
             ("Armando Gaytan Godinez", "Asesor responsable"), None]
    for k, d in enumerate(datos):
        cel = t.rows[k // 2].cells[k % 2]; cel.text = ""
        if not d:
            continue
        for n, (txt, neg, tam, antes) in enumerate((("", False, 12, 40), ("_" * 30, False, 12, 0),
                                                     (d[0], True, 11, 2), (d[1], False, 10, 0))):
            p = cel.paragraphs[0] if n == 0 else cel.add_paragraph()
            base.formato(p, WD_ALIGN_PARAGRAPH.CENTER, 1.0, antes, 0)
            base.texto(p, txt, tam, negrita=neg)


# ---------------------------------------------------------------- documento
def construir(paginas):
    md = open(MD, encoding="utf-8").read()
    subtitulo = re.search(r"^###\s+(.*)$", md, re.M).group(1)
    doc = Document(base.PLANTILLA)
    normal = doc.styles.default(WD_STYLE_TYPE.PARAGRAPH)
    normal.font.name = "Arial"; normal.font.size = Pt(12); normal.paragraph_format.line_spacing = 1.5
    vinetas, numeradas = base.numeraciones(doc)
    ORDEN_FIG.clear()

    def c(t, tam=12, negrita=True, cursiva=False, antes=6, despues=4, alin=WD_ALIGN_PARAGRAPH.CENTER):
        p = doc.add_paragraph(); base.formato(p, alin, 1.15, antes, despues)
        base.texto(p, t, tam, negrita=negrita, cursiva=cursiva)

    # Caratula
    c("Tecnológico Nacional de México", 13, antes=0, despues=0)
    c("Instituto Tecnológico de Colima", 13, antes=0, despues=10)
    c("REPORTE TÉCNICO FINAL", 16, cursiva=True, antes=4, despues=4)
    c(subtitulo, 13, cursiva=True, antes=2, despues=8)
    p = doc.add_paragraph(); base.formato(p, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 4, 4)
    p.add_run().add_picture(os.path.join(base.EVID, "01-mapa-casa-luis.png"), width=Cm(8.0))
    c("Ingeniería Mecatrónica. Servicio Social en Investigación Tecnológica.", 12, cursiva=True, antes=4, despues=10)
    c("Asesor responsable:", 12, alin=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=8, despues=0)
    c("Armando Gaytan Godinez.", 12, cursiva=True, alin=WD_ALIGN_PARAGRAPH.LEFT, antes=4, despues=6)
    c("Alumnos:", 12, alin=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=8, despues=0)
    for nombre, control in (("Flores Bueno Luis Adrian", "22460736"), ("Toscano Farias Andros Jair", "22460548")):
        c(nombre + ".  No. de control: " + control, 12, alin=WD_ALIGN_PARAGRAPH.LEFT, antes=6, despues=0)
    c("Periodo:", 12, alin=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=12, despues=0)
    c(base.PERIODO + ".", 12, cursiva=True, alin=WD_ALIGN_PARAGRAPH.LEFT, antes=4, despues=0)
    c(base.FECHA, 12, alin=WD_ALIGN_PARAGRAPH.RIGHT, antes=14, despues=0)
    base.salto(doc)

    # Indice (se rellena con los titulos de la conversion: se reserva su sitio)
    base.titulo(doc, "Índice.", 1)
    marca_indice = doc.add_paragraph()
    base.salto(doc)
    base.titulo(doc, "Tabla de ilustraciones.", 1)
    marca_figuras = doc.add_paragraph()
    base.salto(doc)

    titulos = convertir(doc, md, vinetas, numeradas)

    # Indice
    anclas = []
    def entrada(t, nivel):
        doc_tmp = doc.add_paragraph()   # se crea al final y se mueve a su sitio
        base.formato(doc_tmp, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 4 if nivel == 1 else 1, 0)
        doc_tmp.paragraph_format.left_indent = Cm(0 if nivel == 1 else 0.9)
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        doc_tmp.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        tam = 11 if nivel == 1 else 10
        base.texto(doc_tmp, t, tam, negrita=True, cursiva=True)
        if paginas.get(t):
            base.fuente(doc_tmp.add_run("\t" + str(paginas[t])), tam, True, True)
        anclas.append(doc_tmp)
    entrada("Tabla de ilustraciones.", 1)
    for nivel, t in titulos:
        entrada(t, nivel)
    for par in anclas:
        marca_indice._element.addprevious(par._element)
    marca_indice._element.getparent().remove(marca_indice._element)

    # Tabla de ilustraciones con miniaturas, como la referencia
    filas = [["Fig %d." % (k + 1), d, str(paginas.get("Fig %d." % (k + 1), ""))] for k, (_r, d) in enumerate(ORDEN_FIG)]
    antes_n = len(doc.element.body)
    base.tabla(doc, ["Código", "Descripción de la ilustración / figura", "Página"], filas, [4.2, 8.8, 2.5],
               miniaturas=[r for r, _d in ORDEN_FIG], alto_miniatura=1.25)
    nuevos = list(doc.element.body)[antes_n - 1:len(doc.element.body) - 1]
    for el in nuevos:
        if el.tag != qn("w:sectPr"):
            marca_figuras._element.addprevious(el)
    marca_figuras._element.getparent().remove(marca_figuras._element)
    return doc, titulos


def main():
    paginas = json.load(open(sys.argv[1], encoding="utf-8")) if len(sys.argv) > 1 and os.path.exists(sys.argv[1]) else {}
    doc, titulos = construir(paginas)
    doc.save(SALIDA)
    json.dump({"titulos": ["Tabla de ilustraciones."] + [t for n, t in titulos if n == 1 and t != "Firmas."],
               "subtitulos": [t for n, t in titulos if n == 2],
               "figuras": ["Fig %d." % (k + 1) for k in range(len(ORDEN_FIG))],
               "firmas": "Firmas."}, open(os.path.join(AQUI, "titulos_tecnico.json"), "w", encoding="utf-8"),
              ensure_ascii=False)
    print("Escrito", os.path.relpath(SALIDA, RAIZ), "-", len(titulos), "titulos,", len(ORDEN_FIG), "figuras")


if __name__ == "__main__":
    main()
